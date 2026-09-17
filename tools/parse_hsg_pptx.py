# -*- coding: utf-8 -*-
"""
parse_hsg_pptx.py
Parse tất cả PPTX HSG từ 2 nguồn:
  - D:\Teacher OS\HSG\PPT_On_Thi_HSG\*.pptx  (tự nhiên)
  - D:\Teacher OS\Tổ KHXH\Giao_an_26_27\HSG 9\*.pptx  (vùng KT)
Output: hsg_slider\extracted\<ten_file>.json cho mỗi PPTX
"""
import sys
import os
import json
import re
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("ERROR: python-pptx not installed. Run: pip install python-pptx")
    sys.exit(1)

# Thư mục nguồn
SOURCES = [
    Path(r"D:\Teacher OS\HSG\PPT_On_Thi_HSG"),
    Path(r"D:\Teacher OS\Tổ KHXH\Giao_an_26_27\HSG 9"),
]

# Thư mục output (workspace-relative)
SCRIPT_DIR = Path(__file__).parent
PROJECT_DIR = SCRIPT_DIR.parent
OUTPUT_DIR = PROJECT_DIR / "hsg_slider" / "extracted"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def extract_text_from_shape(shape):
    """Trích text từ shape (text frame)."""
    texts = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = " ".join(run.text for run in para.runs if run.text.strip())
            if line.strip():
                texts.append(line.strip())
    return texts


def extract_table_from_shape(shape):
    """Trích tất cả cell text từ table."""
    rows_data = []
    if shape.has_table:
        for row in shape.table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                rows_data.append(" | ".join(row_texts))
    return rows_data


def parse_pptx_file(pptx_path: Path):
    """Parse 1 file PPTX, trả về list các item."""
    items = []
    filename = pptx_path.stem

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        print(f"  ERROR opening {pptx_path.name}: {e}")
        return items

    for slide_num, slide in enumerate(prs.slides, start=1):
        # 1. Text frames từ shapes
        for shape in slide.shapes:
            # Text thường
            shape_texts = extract_text_from_shape(shape)
            for txt in shape_texts:
                if len(txt) > 3:  # Bỏ qua ký tự rỗng/quá ngắn
                    items.append({
                        "file": pptx_path.name,
                        "slide": slide_num,
                        "source": "shape",
                        "text": txt
                    })

            # Table
            table_texts = extract_table_from_shape(shape)
            for txt in table_texts:
                if len(txt) > 3:
                    items.append({
                        "file": pptx_path.name,
                        "slide": slide_num,
                        "source": "table",
                        "text": txt
                    })

            # Group shapes (nested)
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP = 6
                try:
                    for sub_shape in shape.shapes:
                        for txt in extract_text_from_shape(sub_shape):
                            if len(txt) > 3:
                                items.append({
                                    "file": pptx_path.name,
                                    "slide": slide_num,
                                    "source": "shape_group",
                                    "text": txt
                                })
                        for txt in extract_table_from_shape(sub_shape):
                            if len(txt) > 3:
                                items.append({
                                    "file": pptx_path.name,
                                    "slide": slide_num,
                                    "source": "table_group",
                                    "text": txt
                                })
                except Exception:
                    pass

        # 2. Speaker notes
        try:
            notes_slide = slide.notes_slide
            notes_text = notes_slide.notes_text_frame.text.strip()
            if notes_text and len(notes_text) > 3:
                items.append({
                    "file": pptx_path.name,
                    "slide": slide_num,
                    "source": "notes",
                    "text": notes_text
                })
        except Exception:
            pass

    return items


def main():
    all_files = []
    for src_dir in SOURCES:
        if not src_dir.exists():
            print(f"WARNING: Source dir not found: {src_dir}")
            continue
        for f in sorted(src_dir.iterdir()):
            if f.suffix.lower() in ('.pptx', '.ppt') and not f.name.startswith('~'):
                all_files.append(f)

    print(f"Found {len(all_files)} PPTX/PPT files")

    total_items = 0
    for pptx_path in all_files:
        print(f"  Parsing: {pptx_path.name} ...")
        items = parse_pptx_file(pptx_path)
        print(f"    -> {len(items)} items extracted")

        # Lưu ra file JSON
        safe_name = re.sub(r'[^\w\-_]', '_', pptx_path.stem)
        out_path = OUTPUT_DIR / f"{safe_name}.json"
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(items, f, ensure_ascii=False, indent=2)

        total_items += len(items)

    print(f"\nDone! Total items: {total_items}")
    print(f"Output dir: {OUTPUT_DIR}")

    # Tạo file index
    index = {
        "total_files": len(all_files),
        "total_items": total_items,
        "files": [re.sub(r'[^\w\-_]', '_', f.stem) + ".json" for f in all_files]
    }
    with open(OUTPUT_DIR / "_index.json", 'w', encoding='utf-8') as f:
        json.dump(index, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
